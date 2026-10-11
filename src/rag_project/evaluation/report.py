from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from rag_project.evaluation.runner import EvalSample


# ============================================================
# Styling
# ============================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

TITLE_FONT = Font(
    bold=True,
    size=16,
)

SUBTITLE_FONT = Font(
    bold=True,
    size=11,
)

THIN_BORDER = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2"),
)


# ============================================================
# Helpers
# ============================================================

def style_header(ws):
    """Style the first row of a worksheet."""

    for cell in ws[1]:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    ws.freeze_panes = "A2"


def auto_width(ws, max_width=60):
    """Automatically size columns."""

    for column_cells in ws.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0

        for cell in column_cells:

            value = cell.value

            if value is None:
                continue

            max_length = max(
                max_length,
                len(str(value)),
            )

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            max_width,
        )


def add_title(ws, title, subtitle=None):
    """Add a title above a worksheet."""

    ws.insert_rows(1, 2)

    ws["A1"] = title
    ws["A1"].font = TITLE_FONT

    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = SUBTITLE_FONT


def format_data_sheet(ws):
    """Apply common formatting."""

    style_header(ws)

    for row in ws.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            cell.border = THIN_BORDER

    auto_width(ws)


# ============================================================
# RAGAS score helpers
# ============================================================

def get_score(df, column):
    """
    Return the mean of a RAGAS score column.

    Returns None if the column does not exist.
    """

    if column not in df.columns:
        return None

    values = pd.to_numeric(
        df[column],
        errors="coerce",
    )

    if values.dropna().empty:
        return None

    return float(values.mean())


# ============================================================
# Build Questions dataframe
# ============================================================

def build_questions_dataframe(
    samples: list[EvalSample],
    ragas_df: pd.DataFrame,
) -> pd.DataFrame:

    rows = []

    # RAGAS rows correspond to successful samples.
    valid_samples = [
        sample
        for sample in samples
        if sample.error is None
    ]

    for index, sample in enumerate(valid_samples):

        row = {
            "ID": sample.id,
            "Language": sample.language.upper(),
            "Question": sample.question,
            "Expected Article": sample.expected_article,
            "Retrieved Articles": ", ".join(
                str(x)
                for x in sample.retrieved_ids
            ),
            "Article Hit": (
                sample.expected_article
                in sample.retrieved_ids
                if sample.expected_article is not None
                else False
            ),
            "Answer": sample.answer,
            "Reference Answer": sample.reference,
        }

        # Add RAGAS scores if available.
        if index < len(ragas_df):

            ragas_row = ragas_df.iloc[index]

            for column in [
                "faithfulness",
                "llm_context_recall",
                "llm_context_precision_with_reference",
                "factual_correctness",
            ]:

                if column in ragas_df.columns:

                    row[column] = ragas_row[column]

        rows.append(row)

    return pd.DataFrame(rows)


# ============================================================
# Retrieval dataframe
# ============================================================

def build_retrieval_dataframe(
    samples: list[EvalSample],
) -> pd.DataFrame:

    rows = []

    for sample in samples:

        if sample.error:
            continue

        expected = sample.expected_article

        retrieved = sample.retrieved_ids

        hit = (
            expected in retrieved
            if expected is not None
            else False
        )

        rows.append(
            {
                "ID": sample.id,
                "Language": sample.language.upper(),
                "Expected Article": expected,
                "Retrieved Articles": ", ".join(
                    str(x)
                    for x in retrieved
                ),
                "Number Retrieved": len(retrieved),
                "Article Hit": hit,
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# Errors dataframe
# ============================================================

def build_errors_dataframe(
    samples: list[EvalSample],
) -> pd.DataFrame:

    rows = []

    for sample in samples:

        if not sample.error:
            continue

        rows.append(
            {
                "ID": sample.id,
                "Language": sample.language.upper(),
                "Question": sample.question,
                "Error": sample.error,
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# Create Excel report
# ============================================================

def create_excel_report(
    samples: list[EvalSample],
    ragas_result,
    output_path: Path,
):
    """
    Create a professional Excel evaluation report.
    """

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    ragas_df = ragas_result.to_pandas()

    questions_df = build_questions_dataframe(
        samples,
        ragas_df,
    )

    retrieval_df = build_retrieval_dataframe(
        samples,
    )

    errors_df = build_errors_dataframe(
        samples,
    )

    # --------------------------------------------------------
    # Workbook
    # --------------------------------------------------------

    wb = Workbook()

    # Remove default sheet.
    default_ws = wb.active
    wb.remove(default_ws)

    # ========================================================
    # Summary
    # ========================================================

    ws = wb.create_sheet("Summary")

    language = (
        samples[0].language.upper()
        if samples
        else "UNKNOWN"
    )

    successful = sum(
        sample.error is None
        for sample in samples
    )

    failed = len(samples) - successful

    article_hits = sum(
        (
            sample.expected_article
            in sample.retrieved_ids
        )
        for sample in samples
        if sample.error is None
        and sample.expected_article is not None
    )

    article_hit_total = sum(
        (
            sample.expected_article is not None
        )
        for sample in samples
        if sample.error is None
    )

    article_hit_rate = (
        article_hits / article_hit_total
        if article_hit_total
        else None
    )

    ws["A1"] = "RAG Evaluation Report"
    ws["A1"].font = TITLE_FONT

    ws["A2"] = f"Language: {language}"
    ws["A2"].font = SUBTITLE_FONT

    ws["A4"] = "Evaluation Overview"
    ws["A4"].font = SUBTITLE_FONT

    overview = [
        ("Total Questions", len(samples)),
        ("Successful", successful),
        ("Failed", failed),
        (
            "Article Hit Rate",
            article_hit_rate,
        ),
    ]

    row = 5

    for label, value in overview:

        ws.cell(
            row=row,
            column=1,
            value=label,
        )

        ws.cell(
            row=row,
            column=2,
            value=value,
        )

        row += 1

    # --------------------------------------------------------
    # RAGAS scores
    # --------------------------------------------------------

    ws["D4"] = "RAGAS Metrics"
    ws["D4"].font = SUBTITLE_FONT

    metric_names = [
        (
            "Faithfulness",
            "faithfulness",
        ),
        (
            "Context Recall",
            "llm_context_recall",
        ),
        (
            "Context Precision",
            "llm_context_precision_with_reference",
        ),
        (
            "Factual Correctness",
            "factual_correctness",
        ),
    ]

    row = 5

    for label, column in metric_names:

        ws.cell(
            row=row,
            column=4,
            value=label,
        )

        score = get_score(
            ragas_df,
            column,
        )

        cell = ws.cell(
            row=row,
            column=5,
            value=score,
        )

        if score is not None:
            cell.number_format = "0.00"

        row += 1

    # --------------------------------------------------------
    # Interpretation
    # --------------------------------------------------------

    ws["A12"] = "What the metrics mean"
    ws["A12"].font = SUBTITLE_FONT

    explanations = [
        (
            "Faithfulness",
            "Whether the generated answer is supported by the retrieved context."
        ),
        (
            "Context Recall",
            "Whether the retrieved context contains the information needed for the reference answer."
        ),
        (
            "Context Precision",
            "Whether relevant retrieved information is prioritized over irrelevant information."
        ),
        (
            "Factual Correctness",
            "Whether the generated answer agrees with the reference answer."
        ),
        (
            "Article Hit Rate",
            "Whether the expected legal article appeared in the retrieved article IDs."
        ),
    ]

    row = 13

    for metric, explanation in explanations:

        ws.cell(
            row=row,
            column=1,
            value=metric,
        ).font = Font(bold=True)

        ws.merge_cells(
            start_row=row,
            start_column=2,
            end_row=row,
            end_column=5,
        )

        ws.cell(
            row=row,
            column=2,
            value=explanation,
        ).alignment = Alignment(
            wrap_text=True
        )

        row += 1

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 5
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 20

    # ========================================================
    # Questions
    # ========================================================

    ws = wb.create_sheet("Questions")

    if questions_df.empty:

        ws["A1"] = "No successful questions."

    else:

        for col_index, column in enumerate(
            questions_df.columns,
            start=1,
        ):

            ws.cell(
                row=1,
                column=col_index,
                value=column,
            )

        for row_index, values in enumerate(
            questions_df.itertuples(index=False),
            start=2,
        ):

            for col_index, value in enumerate(
                values,
                start=1,
            ):

                ws.cell(
                    row=row_index,
                    column=col_index,
                    value=value,
                )

        format_data_sheet(ws)

    # ========================================================
    # Retrieval
    # ========================================================

    ws = wb.create_sheet("Retrieval")

    if retrieval_df.empty:

        ws["A1"] = "No retrieval results."

    else:

        for col_index, column in enumerate(
            retrieval_df.columns,
            start=1,
        ):

            ws.cell(
                row=1,
                column=col_index,
                value=column,
            )

        for row_index, values in enumerate(
            retrieval_df.itertuples(index=False),
            start=2,
        ):

            for col_index, value in enumerate(
                values,
                start=1,
            ):

                ws.cell(
                    row=row_index,
                    column=col_index,
                    value=value,
                )

        format_data_sheet(ws)

    # ========================================================
    # Errors
    # ========================================================

    ws = wb.create_sheet("Errors")

    if errors_df.empty:

        ws["A1"] = "No errors."

    else:

        for col_index, column in enumerate(
            errors_df.columns,
            start=1,
        ):

            ws.cell(
                row=1,
                column=col_index,
                value=column,
            )

        for row_index, values in enumerate(
            errors_df.itertuples(index=False),
            start=2,
        ):

            for col_index, value in enumerate(
                values,
                start=1,
            ):

                ws.cell(
                    row=row_index,
                    column=col_index,
                    value=value,
                )

        format_data_sheet(ws)

    # ========================================================
    # Save
    # ========================================================

    wb.save(output_path)

    print(
        f"\nReport saved to:\n{output_path}"
    )